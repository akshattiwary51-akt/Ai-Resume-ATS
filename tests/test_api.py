"""Tests for FastAPI endpoints."""
import os
import pytest
from fastapi.testclient import TestClient

from backend.main import app
from backend.core.config import APP_VERSION

@pytest.fixture
def client():
    # Use lifespan context if needed or basic TestClient
    with TestClient(app) as test_client:
        yield test_client

def test_root_endpoint(client):
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert data["name"] == "ATS Resume Analyzer API"
    assert data["version"] == APP_VERSION
    assert "POST   /api/v1/analyze-resume" in data["endpoints"]

def test_health_endpoint(client):
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"

def test_analyze_resume_endpoint_txt_dev_mode(client):
    # Tests resume analysis in dev mode with mock token
    sample_resume = (
        "John Doe\njohn@example.com\n\n"
        "Summary\nExperienced developer.\n\n"
        "Skills\nPython, Docker, SQL\n\n"
        "Experience\nSoftware Engineer at TechCo\n"
        "- Developed backend microservices using Python and Docker.\n"
    )
    files = {
        "resume": ("resume.txt", sample_resume.encode("utf-8"), "text/plain"),
    }
    data = {
        "job_description": "We are seeking a Python developer with Docker skills.",
    }
    headers = {
        "Authorization": "Bearer dev-token",
    }
    response = client.post("/api/v1/analyze-resume", files=files, data=data, headers=headers)
    assert response.status_code == 200
    res = response.json()
    assert "ATS_score" in res
    assert 0 <= res["ATS_score"] <= 100
    assert "component_scores" in res
    assert "detailed_feedback" in res


def test_rewrite_bullet_api_endpoint(client):
    headers = {"Authorization": "Bearer dev-token"}
    payload = {
        "bullet": "helped with docker containers and kubernetes deployments",
        "target_role": "DevOps Engineer",
        "target_skill": "Kubernetes",
    }
    response = client.post("/api/v1/rewrite-bullet", json=payload, headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert "rewritten" in data
    assert "action_verb" in data
    assert len(data["suggested_metrics"]) >= 1


def test_cover_letter_api_endpoint(client):
    headers = {"Authorization": "Bearer dev-token"}
    payload = {
        "candidate_name": "Jordan Lee",
        "skills": ["Python", "FastAPI"],
        "experience_highlights": ["built scalable microservices"],
        "job_title": "Backend Engineer",
        "company_name": "Tech Corp",
    }
    response = client.post("/api/v1/generate-cover-letter", json=payload, headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert "cover_letter" in data
    assert "Jordan Lee" in data["cover_letter"]
    assert "Tech Corp" in data["cover_letter"]


def test_interview_prep_api_endpoint(client):
    headers = {"Authorization": "Bearer dev-token"}
    payload = {
        "skills": ["Python", "Docker"],
        "job_title": "Platform Engineer",
        "unvalidated_skills": ["Redis"],
    }
    response = client.post("/api/v1/interview-prep", json=payload, headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert data["target_role"] == "Platform Engineer"
    assert len(data["technical_questions"]) >= 1


def test_skill_taxonomy_api_endpoint(client):
    response = client.get("/api/v1/skill-taxonomy")
    assert response.status_code == 200
    data = response.json()
    assert data["total_skills"] > 20
    assert any(s["canonical"] == "Kubernetes" for s in data["skills"])


def test_analyze_resume_endpoint_pdf_success(client):
    headers = {"Authorization": "Bearer dev-token"}
    pdf_bytes = b"%PDF-1.4 synthetic valid pdf bytes header"
    extracted_text = (
        "Jane Doe\njane@example.com\n\n"
        "Summary\nSenior software engineer with 5 years experience in building APIs.\n\n"
        "Skills\nPython, FastAPI, Docker, SQL\n\n"
        "Experience\nLead Developer at Acme Corp\n- Built microservices using Python and Docker.\n"
    )

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr("backend.services.resume_parser.extract_text_from_pdf", lambda data: extracted_text)
        files = {"resume": ("resume.pdf", pdf_bytes, "application/pdf")}
        response = client.post("/api/v1/analyze-resume", files=files, headers=headers)
        assert response.status_code == 200
        res = response.json()
        assert "ATS_score" in res
        assert 0 <= res["ATS_score"] <= 100


def test_analyze_resume_endpoint_docx_success(client):
    headers = {"Authorization": "Bearer dev-token"}
    from docx import Document
    import io
    doc = Document()
    doc.add_paragraph("Alice Smith\nalice@example.com")
    doc.add_paragraph("Summary: Experienced Python developer.")
    doc.add_paragraph("Skills: Python, React, Docker")
    doc.add_paragraph("Experience: Software Engineer building backend systems.")
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    files = {"resume": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/v1/analyze-resume", files=files, headers=headers)
    assert response.status_code == 200
    res = response.json()
    assert "ATS_score" in res


def test_analyze_resume_endpoint_corrupted_pdf(client):
    headers = {"Authorization": "Bearer dev-token"}
    corrupt_bytes = b"garbage text that is definitely not a pdf"
    files = {"resume": ("resume.pdf", corrupt_bytes, "application/pdf")}
    response = client.post("/api/v1/analyze-resume", files=files, headers=headers)
    assert response.status_code in (400, 422)
    data = response.json()
    assert "detail" in data
    assert any(term in data["detail"].lower() for term in ("pdf", "corrupted", "format", "valid"))


def test_analyze_resume_endpoint_textless_pdf(client):
    headers = {"Authorization": "Bearer dev-token"}
    from pypdf import PdfWriter
    import io
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    bio = io.BytesIO()
    writer.write(bio)
    textless_pdf = bio.getvalue()

    files = {"resume": ("resume.pdf", textless_pdf, "application/pdf")}
    response = client.post("/api/v1/analyze-resume", files=files, headers=headers)
    assert response.status_code == 422
    data = response.json()
    assert "selectable text" in data["detail"].lower() or "scanned" in data["detail"].lower()


def test_analyze_resume_endpoint_unsupported_extension(client):
    headers = {"Authorization": "Bearer dev-token"}
    files = {"resume": ("resume.exe", b"binary content", "application/octet-stream")}
    response = client.post("/api/v1/analyze-resume", files=files, headers=headers)
    assert response.status_code == 400
    data = response.json()
    assert "unsupported file extension" in data["detail"].lower()


